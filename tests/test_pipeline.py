"""Tests for the housing research agent.

Everything here runs offline: no API key, no web access.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path
from types import SimpleNamespace

import numpy as np
import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from sample_data.generate import build, write_sample, write_sql_sample  # noqa: E402
from uzhousing import pipeline  # noqa: E402
from uzhousing.analysis import drivers as drivers_mod  # noqa: E402
from uzhousing.analysis import metrics, regional, timeseries  # noqa: E402
from uzhousing.config import Settings  # noqa: E402
from uzhousing.ingest import loader, profiler, relevance  # noqa: E402
from uzhousing import llm as llm_mod  # noqa: E402
from uzhousing.llm import extract_json  # noqa: E402
from uzhousing.report import narrative as narrative_mod  # noqa: E402


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------
@pytest.fixture(scope="session")
def sample_json(tmp_path_factory) -> Path:
    path = tmp_path_factory.mktemp("data") / "sample.json"
    return write_sample(path)


@pytest.fixture(scope="session")
def sample_sql(tmp_path_factory) -> Path:
    path = tmp_path_factory.mktemp("data") / "sample.sql"
    return write_sql_sample(path)


@pytest.fixture(scope="session")
def understanding(sample_json):
    return profiler.understand(loader.load(sample_json), None)


@pytest.fixture(scope="session")
def analysis(understanding):
    return pipeline.analyse(understanding)


# ---------------------------------------------------------------------------
# Ingestion
# ---------------------------------------------------------------------------
class TestLoader:
    def test_loads_multi_table_json(self, sample_json):
        dataset = loader.load(sample_json)
        assert set(dataset.tables) == {"regional_housing", "macro_indicators"}
        assert dataset.total_rows > 1000

    def test_loads_sql_script(self, sample_sql):
        dataset = loader.load(sample_sql)
        assert "regional_housing" in dataset.tables
        assert len(dataset.tables["regional_housing"]) > 1000

    def test_json_and_sql_agree(self, sample_json, sample_sql):
        from_json = loader.load(sample_json).tables["regional_housing"]
        from_sql = loader.load(sample_sql).tables["regional_housing"]
        assert len(from_json) == len(from_sql)
        assert np.isclose(
            from_json["avg_price_per_sqm_usd"].sum(),
            from_sql["avg_price_per_sqm_usd"].sum(),
        )

    def test_list_of_records(self, tmp_path):
        path = tmp_path / "flat.json"
        path.write_text(json.dumps([
            {"date": "2024-01", "price": 100},
            {"date": "2024-02", "price": 110},
        ]))
        assert len(loader.load(path).tables["flat"]) == 2

    def test_payload_wrapper_and_metadata(self, tmp_path):
        path = tmp_path / "wrapped.json"
        path.write_text(json.dumps({
            "meta": "monthly",
            "data": [{"date": "2024-01", "price": 100}, {"date": "2024-02", "price": 120}],
        }))
        dataset = loader.load(path)
        assert dataset.total_rows == 2

    def test_column_oriented_json(self, tmp_path):
        path = tmp_path / "columns.json"
        path.write_text(json.dumps({"date": ["2024-01", "2024-02"], "price": [1, 2]}))
        frame = next(iter(loader.load(path).tables.values()))
        assert list(frame.columns) == ["date", "price"]

    def test_messy_numeric_text_is_coerced(self, tmp_path):
        path = tmp_path / "messy.csv"
        path.write_text("date,price,rate\n2024-01,\"1 234,5\",12%\n2024-02,\"1 300,0\",11%\n")
        frame = next(iter(loader.load(path).tables.values()))
        assert frame["price"].dtype.kind == "f"
        assert frame["price"].iloc[0] == pytest.approx(1234.5)
        assert frame["rate"].iloc[0] == pytest.approx(12.0)

    def test_missing_file_raises(self, tmp_path):
        with pytest.raises(loader.LoadError):
            loader.load(tmp_path / "nope.json")

    def test_json_arrays_in_a_field_do_not_break_profiling(self, tmp_path):
        """Regression: list cells are unhashable and used to crash nunique()."""
        path = tmp_path / "nested.json"
        path.write_text(json.dumps([
            {
                "date": f"2024-{m:02d}",
                "region": "Tashkent",
                "avg_price_per_sqm_usd": 1000 + m * 10,
                "transactions": 500 + m,
                "tags": ["primary", "urban"],
                "breakdown": {"rooms": [1, 2, 3]},
                "empty": [],
            }
            for m in range(1, 13)
        ]))

        dataset = loader.load(path)
        frame = next(iter(dataset.tables.values()))
        assert not frame["tags"].map(lambda v: isinstance(v, list)).any()
        assert frame["tags"].iloc[0] == '["primary", "urban"]'
        assert any("lists or nested objects" in note for note in dataset.notes)

        # The whole understanding step must survive it.
        result = profiler.understand(dataset, None)
        assert "avg_price_per_sqm_usd" in result.metrics
        assert result.primary_profile.date_col == "date"

    def test_sets_and_tuples_are_encoded_too(self, tmp_path):
        path = tmp_path / "odd.csv"
        path.write_text("a,b\n1,2\n3,4\n")
        frame = next(iter(loader.load(path).tables.values()))
        frame["c"] = [(1, 2), {"x"}]
        assert loader._flatten_containers(frame) == ["c"]
        assert frame["c"].iloc[0] == "[1, 2]"

    def test_parallel_arrays_are_expanded_into_rows(self, tmp_path):
        path = tmp_path / "parallel.json"
        path.write_text(json.dumps([
            {
                "region": "Tashkent",
                "date": ["2024-01", "2024-02", "2024-03"],
                "avg_price_per_sqm_usd": [1000, 1010, 1025],
            },
            {
                "region": "Samarkand",
                "date": ["2024-01", "2024-02", "2024-03"],
                "avg_price_per_sqm_usd": [700, 705, 715],
            },
        ]))

        dataset = loader.load(path)
        frame = next(iter(dataset.tables.values()))
        assert len(frame) == 6
        assert frame["avg_price_per_sqm_usd"].dtype.kind in "if"
        assert set(frame["region"]) == {"Tashkent", "Samarkand"}
        assert any("parallel arrays" in note for note in dataset.notes)

        result = profiler.understand(dataset, None)
        assert "avg_price_per_sqm_usd" in result.metrics
        assert result.regions == ["Samarkand", "Tashkent"]

    def test_a_single_list_column_is_not_exploded(self, tmp_path):
        """Exploding one array alone would duplicate and double-count the measures."""
        path = tmp_path / "tags.json"
        path.write_text(json.dumps([
            {"date": "2024-01", "price": 100, "tags": ["a", "b"]},
            {"date": "2024-02", "price": 110, "tags": ["c"]},
        ]))
        frame = next(iter(loader.load(path).tables.values()))
        assert len(frame) == 2
        assert frame["price"].sum() == 210

    def test_ragged_arrays_are_left_alone(self, tmp_path):
        path = tmp_path / "ragged.json"
        path.write_text(json.dumps([
            {"region": "A", "date": ["2024-01", "2024-02"], "price": [1, 2, 3]},
        ]))
        frame = next(iter(loader.load(path).tables.values()))
        assert len(frame) == 1  # stringified rather than mis-expanded

    def test_safe_nunique_handles_unhashable_values(self):
        series = pd.Series([["a"], ["a"], ["b"]], name="tags")
        with pytest.raises(TypeError):
            series.nunique()
        assert profiler.safe_nunique(series) == 2
        assert len(profiler.safe_unique(series)) == 2

    def test_sql_split_ignores_semicolons_in_strings(self):
        statements = loader._split_statements("INSERT INTO t VALUES ('a;b'); SELECT 1;")
        assert len(statements) == 2
        assert "'a;b'" in statements[0]


# ---------------------------------------------------------------------------
# Folders
# ---------------------------------------------------------------------------
def _write_quarter_db(path: Path, year: int, rows: int, start: int = 0) -> Path:
    """A quarterly SQLite dump shaped like the OLX exports."""
    import sqlite3

    conn = sqlite3.connect(path)
    conn.execute("CREATE TABLE prices (date TEXT, region TEXT, price REAL)")
    conn.executemany(
        "INSERT INTO prices VALUES (?, ?, ?)",
        [
            (f"{year}-{1 + i % 12:02d}", "Tashkent", 1000.0 + start + i)
            for i in range(rows)
        ],
    )
    conn.commit()
    conn.close()
    return path


class TestDirectoryLoading:
    def test_reads_every_file_in_a_folder(self, tmp_path):
        folder = tmp_path / "quarters"
        folder.mkdir()
        _write_quarter_db(folder / "prices_2023.db", 2023, 40)
        _write_quarter_db(folder / "prices_2024.db", 2024, 60, start=1000)

        dataset = loader.load(folder)
        assert dataset.total_rows == 100
        assert list(dataset.tables) == ["prices"]  # same shape, stacked into one
        assert dataset.source == str(folder)
        assert any("2 of 2 data files" in note for note in dataset.notes)

    def test_mixed_formats_are_combined(self, tmp_path):
        folder = tmp_path / "mixed"
        folder.mkdir()
        _write_quarter_db(folder / "prices.db", 2023, 10)
        (folder / "extra.csv").write_text("date,region,price\n2024-01,Andijan,900\n")

        dataset = loader.load(folder)
        assert dataset.total_rows == 11
        assert set(dataset.tables) == {"prices", "extra"}

    def test_unreadable_files_are_reported_not_fatal(self, tmp_path):
        folder = tmp_path / "partly_broken"
        folder.mkdir()
        _write_quarter_db(folder / "good.db", 2023, 10)
        (folder / "broken.json").write_text("{not json at all")

        dataset = loader.load(folder)
        assert dataset.total_rows == 10
        assert any("could not be read" in note and "broken.json" in note
                   for note in dataset.notes)

    def test_empty_folder_explains_itself(self, tmp_path):
        folder = tmp_path / "nothing"
        folder.mkdir()
        (folder / "notes.pdf").write_text("not data")

        with pytest.raises(loader.LoadError) as excinfo:
            loader.load(folder)
        assert "no supported data files" in str(excinfo.value)
        assert ".pdf" in str(excinfo.value)

    def test_row_budget_thins_evenly_and_says_so(self, tmp_path):
        """Millions of adverts will not fit in memory, so the loader samples."""
        folder = tmp_path / "big"
        folder.mkdir()
        _write_quarter_db(folder / "a.db", 2023, 500)
        _write_quarter_db(folder / "b.db", 2024, 500, start=5000)

        dataset = loader.load(folder, max_rows=100)
        assert 0 < dataset.total_rows <= 100
        assert dataset.source_rows == 1000

        frame = dataset.tables["prices"]
        assert set(frame["date"].str[:4]) == {"2023", "2024"}  # both files survive
        note = dataset.notes[0]
        assert "1,000 rows" in note and "--max-rows 0" in note

    def test_no_budget_reads_everything(self, tmp_path):
        folder = tmp_path / "small"
        folder.mkdir()
        _write_quarter_db(folder / "a.db", 2023, 250)

        dataset = loader.load(folder, max_rows=0)
        assert dataset.total_rows == 250
        assert not any("--max-rows" in note for note in dataset.notes)

    def test_tables_of_different_shapes_stay_separate(self, tmp_path):
        import sqlite3

        folder = tmp_path / "shapes"
        folder.mkdir()
        _write_quarter_db(folder / "a.db", 2023, 5)
        conn = sqlite3.connect(folder / "b.db")
        conn.execute("CREATE TABLE prices (quite TEXT, different TEXT, columns TEXT)")
        conn.executemany("INSERT INTO prices VALUES (?, ?, ?)", [("x", "y", "z")])
        conn.commit()
        conn.close()

        dataset = loader.load(folder)
        assert len(dataset.tables) == 2
        assert dataset.total_rows == 6


# ---------------------------------------------------------------------------
# Profiling
# ---------------------------------------------------------------------------
class TestProfiler:
    def test_roles_are_detected(self, understanding):
        profile = understanding.primary_profile
        roles = {c.name: c.role for c in profile.columns}
        assert roles["date"] == "date"
        assert roles["region"] == "region"
        assert roles["segment"] == "segment"
        assert roles["avg_price_per_sqm_usd"] == "price_per_sqm"
        assert roles["transactions"] == "volume"
        assert roles["new_units_commissioned"] == "supply"

    def test_grain_and_period(self, understanding):
        profile = understanding.primary_profile
        # The source data is monthly; the agent reports only at quarterly or
        # annual resolution, so it is aggregated up before analysis runs.
        assert profile.grain == "quarterly"
        assert profile.period_start.year == 2018

    def test_supporting_table_metrics_are_merged(self, understanding):
        assert "mortgage_rate_pct" in understanding.metrics
        assert "usd_uzs_rate" in understanding.metrics

    def test_national_placeholder_excluded_from_regions(self, understanding):
        assert "National" not in understanding.regions
        assert "Tashkent City" in understanding.regions

    @pytest.mark.parametrize(
        "raw,expected",
        [
            ("2024Q1", (2024, 1)),
            ("Q3 2021", (2021, 7)),
            ("2020-M07", (2020, 7)),
            ("Mar 2019", (2019, 3)),
            ("2022-11", (2022, 11)),
        ],
    )
    def test_period_formats(self, raw, expected):
        parsed = profiler._parse_special_period(raw.lower()) or pd.to_datetime(raw)
        assert (parsed.year, parsed.month) == expected

    def test_separate_year_and_month_columns(self):
        frame = pd.DataFrame({
            "year": [2023, 2023, 2024],
            "month": [1, 6, 3],
            "price": [10, 11, 12],
        })
        parsed = profiler.parse_dates(frame["year"], frame)
        assert parsed.notna().all()
        assert list(parsed.dt.month) == [1, 6, 3]

    def test_rate_named_column_with_large_values_is_not_a_rate(self):
        series = pd.Series([12000.0, 12500.0, 13000.0], name="usd_uzs_rate")
        role, *_ = profiler.classify_column(series)
        assert role == "fx"

    def test_russian_and_uzbek_names(self):
        for name, expected in [
            ("viloyat", "region"), ("область", "region"),
            ("narx", "price"), ("цена", "price"),
            ("ипотека", "mortgage"), ("stavka", "rate"),
        ]:
            role, *_ = profiler.classify_column(pd.Series(["a", "b"], name=name))
            assert role == expected, f"{name} -> {role}"

    def test_date_names_need_plausible_date_values(self):
        """A date-shaped name is only accepted when the values parse as dates."""
        dates = pd.Series(["2024-01-01", "2024-02-01"], name="sana")
        assert profiler.classify_column(dates)[0] == "date"

        not_dates = pd.Series(["alpha", "beta"], name="sana")
        assert profiler.classify_column(not_dates)[0] != "date"

    def test_mortgage_rate_is_a_rate_not_a_lending_volume(self):
        rate = pd.Series([18.2, 17.9, 17.4], name="mortgage_rate_pct")
        assert profiler.classify_column(rate)[0] == "rate"

        volume = pd.Series([2.1, 2.4, 2.6], name="mortgage_loans_issued_bn_uzs")
        assert profiler.classify_column(volume)[0] == "mortgage"


# ---------------------------------------------------------------------------
# Analysis
# ---------------------------------------------------------------------------
class TestMetrics:
    def test_growth_maths(self):
        index = pd.date_range("2020-01-01", periods=25, freq="MS")
        series = pd.Series(100 * 1.01 ** np.arange(25), index=index)
        summary = metrics.summarise_series(series, "price", role="price")
        assert summary.yoy_pct == pytest.approx(12.68, abs=0.1)
        assert summary.cagr_pct == pytest.approx(12.68, abs=0.3)
        assert summary.direction == "rising"
        assert summary.pct_from_peak == pytest.approx(0.0, abs=1e-6)

    def test_rate_series_reports_percentage_points(self):
        index = pd.date_range("2020-01-01", periods=25, freq="MS")
        series = pd.Series(np.linspace(10, 16, 25), index=index)
        summary = metrics.summarise_series(series, "policy_rate", role="rate")
        assert summary.change_unit == "pp"
        assert summary.yoy_pp == pytest.approx(3.0, abs=0.1)

    def test_empty_series_is_safe(self):
        summary = metrics.summarise_series(pd.Series(dtype=float), "x")
        assert summary.observations == 0
        assert summary.latest is None

    def test_affordability_ratio(self):
        index = pd.date_range("2020-01-01", periods=3, freq="MS")
        price = pd.Series([120000, 130000, 140000], index=index)
        income = pd.Series([1000, 1000, 1000], index=index)
        ratio = metrics.affordability_ratio(price, income)
        assert ratio.iloc[0] == pytest.approx(10.0)


class TestTimeSeries:
    def test_trend_is_detected(self):
        index = pd.date_range("2018-01-01", periods=60, freq="MS")
        series = pd.Series(np.arange(60) * 2.0 + 100, index=index)
        result = timeseries.analyse(series, "price")
        assert result.trend.significant
        assert result.trend.r_squared > 0.99

    def test_structural_break_is_found(self):
        index = pd.date_range("2018-01-01", periods=60, freq="MS")
        values = np.concatenate([np.full(30, 100.0), np.full(30, 150.0)])
        result = timeseries.analyse(pd.Series(values, index=index), "price")
        assert result.breakpoints
        assert result.breakpoints[0].direction == "step up"
        assert pd.Timestamp(result.breakpoints[0].date).year == 2020

    def test_seasonality_is_found(self):
        index = pd.date_range("2018-01-01", periods=72, freq="MS")
        seasonal = 20 * np.sin(2 * np.pi * np.arange(72) / 12)
        result = timeseries.analyse(pd.Series(100 + seasonal, index=index), "volume")
        assert result.seasonality.detected

    def test_forecast_extends_the_index(self):
        index = pd.date_range("2018-01-01", periods=48, freq="MS")
        series = pd.Series(np.arange(48) * 1.5 + 200, index=index)
        result = timeseries.analyse(series, "price")
        assert result.forecast.points
        assert pd.Timestamp(result.forecast.points[0]["date"]) > index[-1]
        for point in result.forecast.points:
            assert point["lower"] <= point["forecast"] <= point["upper"]

    def test_short_series_does_not_crash(self):
        index = pd.date_range("2024-01-01", periods=2, freq="MS")
        result = timeseries.analyse(pd.Series([1.0, 2.0], index=index), "x")
        assert result.forecast.method == "none"


class TestDrivers:
    def _panel(self):
        """A rate that cycles, and a price that responds to it four months later.

        The rate must genuinely cycle: on a year-on-year basis a straight ramp has a
        constant growth rate, so there would be nothing left to correlate.
        """
        index = pd.date_range("2018-01-01", periods=96, freq="MS")
        rng = np.random.default_rng(7)
        t = np.arange(96)
        rate = pd.Series(15 + 4 * np.sin(2 * np.pi * t / 30) + rng.normal(0, 0.15, 96), index=index)
        lagged = np.concatenate([np.full(4, rate.iloc[0]), rate.to_numpy()[:-4]])
        price = pd.Series(900 - 20 * lagged + rng.normal(0, 4, 96), index=index)
        return pd.DataFrame({"price": price, "policy_rate": rate})

    def test_lagged_relationship_is_recovered(self):
        panel = self._panel()
        result = drivers_mod.analyse_drivers(
            panel, "price", {"price": "price", "policy_rate": "rate"}, periods_per_year=12
        )
        assert result.links
        link = result.links[0]
        assert link.driver == "policy_rate"
        assert link.significant
        assert link.correlation < 0  # higher rates, weaker prices

    def test_same_family_driver_is_excluded(self):
        index = pd.date_range("2018-01-01", periods=48, freq="MS")
        frame = pd.DataFrame(
            {
                "price_sqm": np.arange(48, dtype=float) + 100,
                "median_price": np.arange(48, dtype=float) * 60 + 6000,
                "rate": np.linspace(10, 14, 48),
            },
            index=index,
        )
        result = drivers_mod.analyse_drivers(
            frame, "price_sqm",
            {"price_sqm": "price_per_sqm", "median_price": "price", "rate": "rate"},
            periods_per_year=12,
        )
        assert "median_price" in result.excluded
        assert all(link.driver != "median_price" for link in result.links)

    def test_theory_check_only_applies_to_price_targets(self):
        index = pd.date_range("2018-01-01", periods=48, freq="MS")
        frame = pd.DataFrame(
            {
                "transactions": np.arange(48, dtype=float) + 500,
                "completions": np.arange(48, dtype=float) * 2 + 300,
            },
            index=index,
        )
        result = drivers_mod.analyse_drivers(
            frame, "transactions",
            {"transactions": "volume", "completions": "supply"},
            periods_per_year=12,
        )
        # Supply is only expected to move against *prices*, not against volumes.
        assert all(link.consistent_with_theory is None for link in result.links)

    def test_year_on_year_basis_when_history_allows(self):
        index = pd.date_range("2018-01-01", periods=72, freq="MS")
        frame = pd.DataFrame({"a": np.arange(72, dtype=float), "b": np.arange(72, dtype=float)}, index=index)
        result = drivers_mod.analyse_drivers(frame, "a", {"a": "price", "b": "income"}, periods_per_year=12)
        assert result.basis.startswith("year-on-year")
        assert result.basis_note


class TestRegional:
    def test_comparison_ranks_and_measures_spread(self, understanding):
        comparison = regional.compare_groups(
            understanding.tidy, "avg_price_per_sqm_usd", "region", "price_per_sqm"
        )
        assert comparison is not None
        assert comparison.leaders[0]["region"] == "Tashkent City"
        assert comparison.spread_ratio > 1
        assert "diverging" in comparison.convergence

    def test_concentration(self, understanding):
        result = regional.concentration(understanding.tidy, "transactions", "region")
        assert result["top_group"] == "Tashkent City"
        assert 0 < result["top_share_pct"] < 100


# ---------------------------------------------------------------------------
# Report generation
# ---------------------------------------------------------------------------
class TestAnalysisResult:
    def test_headline_is_a_price(self, analysis):
        assert analysis.headline_metric == "avg_price_per_sqm_usd"

    def test_activity_gets_its_own_attribution(self, analysis):
        assert analysis.activity_metric == "transactions"
        assert analysis.activity_drivers is not None

    def test_brief_is_json_serialisable(self, analysis):
        json.dumps(analysis.brief, default=str)

    def test_brief_has_the_expected_sections(self, analysis):
        for key in ("coverage", "headline", "metrics", "timeseries", "regional", "drivers"):
            assert key in analysis.brief, key

    def test_metrics_are_ordered_by_relevance(self, analysis):
        roles = [m["role"] for m in analysis.brief["metrics"]]
        assert roles[0] == "price_per_sqm"
        assert roles.index("area") > roles.index("volume")


class TestRelevanceScreen:
    """Variables are understood first, then screened; the unrelated ones never
    reach a statistic."""

    @staticmethod
    def _understanding() -> profiler.Understanding:
        """A housing panel carrying the junk a scraped marketplace feed brings with it."""
        rows = 36
        frame = pd.DataFrame({
            "date": pd.date_range("2022-01-01", periods=rows, freq="MS"),
            "region": ["Tashkent City"] * rows,
            # Genuine housing and macro measures.
            "avg_price_per_sqm_usd": np.linspace(900, 1300, rows),
            "monthly_rent_usd": np.linspace(220, 340, rows),
            "transactions": np.linspace(400, 700, rows).round(),
            "mortgage_rate_pct": np.linspace(19, 15, rows),
            "floor_area_sqm": np.linspace(58, 62, rows),
            # Marketplace and record-keeping noise.
            "listing_id": np.arange(1, rows + 1),
            "advert_views": np.linspace(50, 5000, rows),
            "photo_count": np.linspace(3, 9, rows).round(),
            "seller_rating": np.linspace(3.5, 4.9, rows),
            "latitude": np.linspace(41.2, 41.4, rows),
            "is_promoted": [0, 1] * (rows // 2),
            "scraped_page_no": np.arange(1, rows + 1),
            "source_version": [3] * rows,
        })
        dataset = loader.Dataset(tables={"market": frame}, source="synthetic", kind="test")
        return profiler.understand(dataset, None)

    def test_marketplace_noise_is_excluded(self):
        understanding = self._understanding()
        screen = relevance.screen(understanding)
        assert screen.applied
        for junk in (
            "advert_views", "photo_count", "seller_rating",
            "latitude", "is_promoted", "scraped_page_no", "source_version",
        ):
            assert junk in screen.dropped, f"{junk} should have been screened out"
        # The profiler already labels a key as an identifier, so it is never a
        # candidate the screen has to reject.
        assert "listing_id" not in understanding.metrics

    def test_a_core_role_does_not_shield_an_obvious_exclusion(self):
        """"photo_count" reads as a count, but counting photographs is not market activity."""
        understanding = self._understanding()
        assert understanding.primary_profile.role_of("photo_count") == "volume"
        screen = relevance.screen(understanding)
        assert "photo_count" in screen.dropped

    def test_housing_measures_survive(self):
        understanding = self._understanding()
        screen = relevance.screen(understanding)
        for kept in (
            "avg_price_per_sqm_usd", "monthly_rent_usd", "transactions",
            "mortgage_rate_pct", "floor_area_sqm",
        ):
            assert kept in screen.kept, f"{kept} should have survived the screen"

    def test_excluded_variables_leave_the_analysis_frame(self):
        understanding = self._understanding()
        screen = relevance.screen(understanding)
        assert "advert_views" not in understanding.metrics
        assert "advert_views" not in understanding.primary_profile.metric_cols
        assert "avg_price_per_sqm_usd" in understanding.metrics
        assert screen.dropped  # the assertions above are not vacuous

    def test_glossary_loses_the_excluded_entries(self):
        understanding = self._understanding()
        glossary = [
            {"name": name, "description": "..."}
            for name in ["avg_price_per_sqm_usd", "advert_views", "seller_rating"]
        ]
        relevance.screen(understanding, glossary)
        assert [entry["name"] for entry in glossary] == ["avg_price_per_sqm_usd"]

    def test_every_decision_carries_a_reason(self):
        understanding = self._understanding()
        screen = relevance.screen(understanding)
        for decision in screen.decisions:
            assert decision.reason
            assert decision.label
            assert decision.source in {"heuristic", "data quality", "model review"}

    def test_a_constant_column_cannot_explain_variation(self):
        col = profiler.ColumnProfile(
            name="price_usd", dtype="float64", role="price",
            non_null=100, missing_pct=0.0, unique=1,
        )
        keep, reason, source = relevance._verdict("price_usd", "price", col)
        assert keep is False
        assert source == "data quality"
        assert "constant" in reason

    def test_screen_is_abandoned_rather_than_emptying_the_dataset(self):
        """Dropping everything is a screening failure, not a finding about the data."""
        rows = 24
        frame = pd.DataFrame({
            "date": pd.date_range("2023-01-01", periods=rows, freq="MS"),
            "advert_views": np.linspace(10, 900, rows),
            "photo_count": np.linspace(1, 9, rows).round(),
        })
        dataset = loader.Dataset(tables={"junk": frame}, source="synthetic", kind="test")
        understanding = profiler.understand(dataset, None)
        screen = relevance.screen(understanding)
        assert not screen.applied
        assert screen.abandoned_reason
        assert not screen.dropped

    def test_the_pipeline_records_the_screen(self, analysis):
        assert analysis.relevance is not None
        assert analysis.relevance.kept
        json.dumps(analysis.relevance.to_dict(), default=str)

    def test_the_screen_stays_out_of_the_report(self, analysis):
        """Decisions are logged for the operator, not narrated to the reader."""
        assert "relevance" not in analysis.brief
        assert "variable_screen" not in analysis.brief
        rendered = json.dumps(analysis.brief, default=str)
        assert "screened out" not in rendered


class TestNarrative:
    def test_fallback_fills_every_section(self, analysis):
        result = narrative_mod.write_fallback(analysis.brief)
        assert result.executive_summary
        assert result.key_findings
        assert result.current_situation
        assert result.historical_trends
        assert result.drivers
        assert result.recommendations
        assert result.risks
        assert result.limitations
        assert result.generated_by == "deterministic template"

    def test_recommendations_are_attributed(self, analysis):
        result = narrative_mod.write_fallback(analysis.brief)
        for recommendation in result.recommendations:
            assert recommendation.audience
            assert recommendation.action
            assert recommendation.rationale

    def test_empty_brief_does_not_crash(self):
        result = narrative_mod.write_fallback({})
        assert result.executive_summary


class TestEndToEnd:
    def test_report_is_produced(self, sample_json, tmp_path):
        settings = Settings.from_env()
        settings.output_dir = tmp_path / "out"
        settings.web_research = False
        settings.anthropic_api_key = ""

        result = pipeline.run(settings=settings, data_path=sample_json, progress=lambda _: None)

        assert result.ok
        assert result.report_path.suffix == ".docx"
        assert result.report_path.stat().st_size > 100_000
        assert len(result.figures) >= 8
        for figure in result.figures:
            assert figure.path.exists()
        assert result.run_log.exists()
        assert json.loads(result.run_log.read_text(encoding="utf-8"))["brief"]

    def test_report_has_headings_figures_and_tables(self, sample_json, tmp_path):
        from docx import Document

        settings = Settings.from_env()
        settings.output_dir = tmp_path / "out2"
        settings.web_research = False
        settings.anthropic_api_key = ""
        result = pipeline.run(settings=settings, data_path=sample_json, progress=lambda _: None)

        doc = Document(str(result.report_path))
        headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
        assert any("Executive summary" in h for h in headings)
        assert any("Recommendations" in h for h in headings)
        assert any("Sources" in h for h in headings)
        assert len(doc.tables) >= 10
        assert len(doc.inline_shapes) >= 8

    def test_sql_source_produces_the_same_headline(self, sample_sql, tmp_path):
        settings = Settings.from_env()
        settings.output_dir = tmp_path / "out3"
        settings.web_research = False
        settings.anthropic_api_key = ""
        result = pipeline.run(settings=settings, data_path=sample_sql, progress=lambda _: None)
        assert result.ok
        assert result.analysis.headline_metric == "avg_price_per_sqm_usd"

    def test_minimal_dataset_still_reports(self, tmp_path):
        """A tiny two-column file should degrade gracefully, not crash."""
        path = tmp_path / "tiny.json"
        rows = [
            {"date": f"2023-{m:02d}", "price": 100 + m * 3}
            for m in range(1, 13)
        ]
        path.write_text(json.dumps(rows))

        settings = Settings.from_env()
        settings.output_dir = tmp_path / "out4"
        settings.web_research = False
        settings.anthropic_api_key = ""
        result = pipeline.run(settings=settings, data_path=path, progress=lambda _: None)
        assert result.ok


# ---------------------------------------------------------------------------
class TestLLMHelpers:
    @pytest.mark.parametrize(
        "raw,expected",
        [
            ('{"a": 1}', {"a": 1}),
            ('```json\n{"a": 1}\n```', {"a": 1}),
            ('Sure!\n{"a": 1}\nHope that helps.', {"a": 1}),
            ("[1, 2, 3]", [1, 2, 3]),
            ("not json at all", None),
            ("", None),
        ],
    )
    def test_json_extraction(self, raw, expected):
        assert extract_json(raw) == expected


class _FakeStream:
    """Stands in for the SDK's streaming context manager."""

    def __init__(self, text: str, stop_reason: str) -> None:
        self._message = SimpleNamespace(
            content=[SimpleNamespace(type="text", text=text)],
            stop_reason=stop_reason,
        )

    def __enter__(self):
        return self

    def __exit__(self, *_exc):
        return False

    def get_final_message(self):
        return self._message


class _FakeClient:
    """Replays canned (text, stop_reason) pairs and records the budget asked for."""

    def __init__(self, replies: list[tuple[str, str]]) -> None:
        self.replies = list(replies)
        self.budgets: list[int] = []
        self.messages = self

    def stream(self, **kwargs):
        self.budgets.append(kwargs["max_tokens"])
        return _FakeStream(*self.replies.pop(0))


def _llm_with(replies: list[tuple[str, str]]) -> tuple[llm_mod.LLM, _FakeClient]:
    client = _FakeClient(replies)
    llm = llm_mod.LLM(api_key="", model="stub-model")
    llm._client = client
    return llm, client


class TestTokenBudget:
    """Adaptive thinking spends the same budget as the answer, so a JSON reply
    can be cut off before its closing brace. That is a budget failure, not a
    model that cannot follow instructions."""

    def test_truncated_json_is_retried_with_a_larger_budget(self):
        llm, client = _llm_with([
            ('{"executive_summary": ["the answer stops mid-sen', "max_tokens"),
            ('{"executive_summary": ["done"]}', "end_turn"),
        ])
        assert llm.complete_json("write the report") == {"executive_summary": ["done"]}
        assert client.budgets[1] > client.budgets[0]

    def test_a_small_budget_is_raised_to_make_room_for_thinking(self):
        llm, client = _llm_with([('{"ok": true}', "end_turn")])
        llm.complete_json("classify this", max_tokens=500)
        assert client.budgets[0] >= llm_mod.JSON_MIN_TOKENS

    def test_budgets_never_exceed_the_ceiling(self):
        llm, client = _llm_with([("{", "max_tokens")] * 6)
        with pytest.raises(llm_mod.LLMUnavailable):
            llm.complete_json("write the report", max_tokens=llm_mod.MAX_OUTPUT_TOKENS)
        assert client.budgets == [llm_mod.MAX_OUTPUT_TOKENS]

    def test_a_complete_but_unparseable_answer_is_not_retried(self):
        """Retrying costs a whole second generation; only truncation earns one."""
        llm, client = _llm_with([("I would rather explain it in prose.", "end_turn")])
        with pytest.raises(llm_mod.LLMUnavailable):
            llm.complete_json("write the report")
        assert len(client.budgets) == 1

    def test_the_stop_reason_reaches_the_caller(self):
        llm, _ = _llm_with([("partial", "max_tokens")])
        text, stop_reason = llm.complete_with_reason("hello")
        assert (text, stop_reason) == ("partial", "max_tokens")


class StubLLM:
    """Stands in for the Anthropic client so the LLM paths can be tested offline.

    Returns canned JSON shaped like the real model's replies, chosen by looking
    for a marker in the prompt.
    """

    model = "stub-model"
    available = True
    status = "stub"

    def __init__(self, *, fail: bool = False) -> None:
        self.fail = fail
        self.calls: list[str] = []

    def complete(self, prompt: str, **_kwargs) -> str:
        raise NotImplementedError

    def complete_json(self, prompt: str, **_kwargs):
        from uzhousing.llm import LLMUnavailable

        self.calls.append(prompt[:60])
        if self.fail:
            raise LLMUnavailable("stub failure")

        if "Allowed roles" in prompt:  # schema review
            return {
                "corrections": [
                    {"table": "regional_housing", "column": "avg_area_sqm",
                     "role": "area", "unit": "m²", "why": "dwelling floor area"}
                ],
                "primary_table": "regional_housing",
                "dataset_summary": "Monthly regional housing prices and activity.",
            }

        if "rigorous, cited synthesis" in prompt:  # research synthesis
            return {
                "themes": [
                    {"key": "policy", "summary": "Policy summary.",
                     "points": ["A specific finding [policy-1]."]}
                ],
                "policy_events": [
                    {"date": "2023-06-01", "title": "Subsidised mortgage expansion",
                     "category": "policy", "summary": "Widened eligibility.",
                     "expected_impact": "Raises demand for primary-market housing.",
                     "direction": "positive", "confidence": "medium", "source_id": "policy-1"}
                ],
                "macro_factors": [
                    {"factor": "Remittances", "current_state": "Rising.",
                     "housing_impact": "Funds deposits.", "direction": "positive",
                     "source_id": "macro-1"}
                ],
            }

        # narrative
        return {
            "executive_summary": ["Summary paragraph one.", "Summary paragraph two."],
            "key_findings": ["Prices rose 4.3% year on year."],
            "current_situation": ["Where the market stands."],
            "historical_trends": ["The shape of the sample."],
            "regional_analysis": ["Tashkent leads."],
            "drivers": ["Rates drive volumes."],
            "policy_analysis": ["The subsidy widened demand."],
            "macro_context": ["Remittances are rising."],
            "outlook": ["Modest growth ahead."],
            "recommendations": [
                {"audience": "Central Bank of Uzbekistan", "action": "Tighten LTV limits.",
                 "rationale": "Price growth outruns incomes.", "priority": "high",
                 "horizon": "6-12 months"}
            ],
            "risks": ["Affordability erosion."],
            "limitations": ["Short sample."],
        }


class TestLLMPaths:
    def test_schema_review_is_applied(self, sample_json):
        dataset = loader.load(sample_json)
        llm = StubLLM()
        result = profiler.understand(dataset, llm)
        assert result.llm_reviewed
        column = result.profiles["regional_housing"].column("avg_area_sqm")
        assert column.role == "area"
        assert "LLM review" in column.reason

    def test_schema_review_failure_falls_back(self, sample_json):
        dataset = loader.load(sample_json)
        result = profiler.understand(dataset, StubLLM(fail=True))
        assert not result.llm_reviewed
        # Heuristics still produced a usable understanding.
        assert result.primary == "regional_housing"
        assert any("skipped" in note or "failed" in note for note in result.notes)

    def test_research_synthesis_is_parsed(self):
        from uzhousing.research import context as research_mod

        findings = research_mod.gather(
            StubLLM(), period_start="2018-01-01", period_end="2026-06-01",
            regions=["Tashkent City"], use_web=False,
            policy_file=ROOT / "knowledge" / "policy_events.json",
            progress=lambda _: None,
        )
        assert findings.llm_used
        assert any(e.title == "Subsidised mortgage expansion" for e in findings.policy_events)
        assert findings.macro_factors[0]["factor"] == "Remittances"
        assert findings.themes[0].key == "policy"

    def test_narrative_uses_the_llm(self, analysis):
        result = narrative_mod.write(analysis.brief, StubLLM(), "en")
        assert result.generated_by.startswith("Claude")
        assert result.executive_summary == ["Summary paragraph one.", "Summary paragraph two."]
        assert result.recommendations[0].audience == "Central Bank of Uzbekistan"
        assert result.recommendations[0].priority == "high"

    def test_narrative_falls_back_when_the_llm_fails(self, analysis):
        result = narrative_mod.write(analysis.brief, StubLLM(fail=True), "en")
        assert result.generated_by == "deterministic template"
        assert result.executive_summary

    def test_full_run_with_llm_stub(self, sample_json, tmp_path, monkeypatch):
        settings = Settings.from_env()
        settings.output_dir = tmp_path / "llm_out"
        settings.web_research = False
        settings.anthropic_api_key = "test-key"

        monkeypatch.setattr(pipeline, "LLM", lambda *_a, **_k: StubLLM())
        result = pipeline.run(settings=settings, data_path=sample_json, progress=lambda _: None)

        assert result.ok
        assert result.narrative.generated_by.startswith("Claude")
        assert result.analysis.understanding.llm_reviewed


class TestSampleData:
    def test_shape(self):
        payload = build()
        assert len(payload["regional_housing"]) == 102 * 8 * 2
        assert len(payload["macro_indicators"]) == 102

    def test_is_deterministic(self):
        assert build()["macro_indicators"] == build()["macro_indicators"]
