"""Render the analysis and narrative into a formatted Word document."""

from __future__ import annotations

import logging
from datetime import date
from pathlib import Path
from typing import Any, Iterable, Sequence

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

LOGGER = logging.getLogger(__name__)

# Ink palette shared with the charts so the document reads as one system.
INK_PRIMARY = RGBColor(0x0B, 0x0B, 0x0B)
INK_SECONDARY = RGBColor(0x52, 0x51, 0x4E)
INK_MUTED = RGBColor(0x89, 0x87, 0x81)
ACCENT = RGBColor(0x2A, 0x78, 0xD6)
ACCENT_DARK = RGBColor(0x18, 0x4F, 0x95)
POSITIVE = RGBColor(0x2A, 0x78, 0xD6)
NEGATIVE = RGBColor(0xE3, 0x49, 0x48)

HEADER_FILL = "2a78d6"
BAND_FILL = "f4f7fd"
RULE_FILL = "e1e0d9"

BODY_FONT = "Calibri"
CONTENT_WIDTH_IN = 6.4


class ReportDocument:
    """Thin, opinionated wrapper over python-docx."""

    def __init__(self, title: str, subtitle: str = "", author: str = "") -> None:
        self.doc = Document()
        self.title = title
        self.subtitle = subtitle
        self.author = author
        self._figure_no = 0
        self._table_no = 0
        self._setup_styles()
        self._setup_page()

    # -- setup ----------------------------------------------------------
    def _setup_page(self) -> None:
        for section in self.doc.sections:
            section.left_margin = Inches(1.0)
            section.right_margin = Inches(1.0)
            section.top_margin = Inches(0.9)
            section.bottom_margin = Inches(0.9)

    def _setup_styles(self) -> None:
        styles = self.doc.styles

        normal = styles["Normal"]
        normal.font.name = BODY_FONT
        normal.font.size = Pt(10.5)
        normal.font.color.rgb = INK_PRIMARY
        normal.paragraph_format.space_after = Pt(8)
        normal.paragraph_format.line_spacing = 1.22
        self._set_east_asian(normal, BODY_FONT)

        for name, size, color, before, after, bold in (
            ("Heading 1", 17, ACCENT_DARK, 22, 8, True),
            ("Heading 2", 13.5, INK_PRIMARY, 16, 6, True),
            ("Heading 3", 11.5, INK_SECONDARY, 12, 4, True),
        ):
            try:
                style = styles[name]
            except KeyError:  # pragma: no cover
                continue
            style.font.name = BODY_FONT
            style.font.size = Pt(size)
            style.font.bold = bold
            style.font.color.rgb = color
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(after)
            style.paragraph_format.keep_with_next = True
            self._set_east_asian(style, BODY_FONT)

    @staticmethod
    def _set_east_asian(style, font_name: str) -> None:
        try:
            style.element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
        except Exception:  # pragma: no cover - style without rPr
            pass

    # -- primitives -----------------------------------------------------
    def heading(self, text: str, level: int = 1) -> None:
        self.doc.add_heading(text, level=level)
        if level == 1:
            self._accent_rule()

    def _accent_rule(self) -> None:
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(10)
        pPr = paragraph._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "10")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), HEADER_FILL)
        borders.append(bottom)
        pPr.append(borders)

    def para(self, text: str, size: float = 10.5, color: RGBColor = INK_PRIMARY,
             italic: bool = False, bold: bool = False, align: str = "justify",
             space_after: float = 8) -> None:
        if not str(text).strip():
            return
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = {
            "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
        }.get(align, WD_ALIGN_PARAGRAPH.JUSTIFY)
        paragraph.paragraph_format.space_after = Pt(space_after)
        run = paragraph.add_run(str(text).strip())
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.italic = italic
        run.bold = bold

    def paragraphs(self, items: Iterable[str]) -> None:
        for item in items:
            self.para(item)

    def bullets(self, items: Iterable[str], style: str = "List Bullet") -> None:
        for item in items:
            if not str(item).strip():
                continue
            try:
                paragraph = self.doc.add_paragraph(str(item).strip(), style=style)
            except KeyError:  # pragma: no cover - template without the style
                paragraph = self.doc.add_paragraph("• " + str(item).strip())
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in paragraph.runs:
                run.font.size = Pt(10.5)
                run.font.color.rgb = INK_PRIMARY

    def callout(self, text: str, label: str = "") -> None:
        """A shaded, single-cell box for a key message."""
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        _shade(cell, BAND_FILL)
        _cell_margins(cell, top=140, bottom=140, left=160, right=160)
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        if label:
            run = paragraph.add_run(label.upper() + "  ")
            run.font.size = Pt(8)
            run.font.bold = True
            run.font.color.rgb = ACCENT
        run = paragraph.add_run(str(text).strip())
        run.font.size = Pt(10)
        run.font.color.rgb = INK_PRIMARY
        _table_borders(table, color=HEADER_FILL, size=4, edges=("left",))
        self.doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -- figures --------------------------------------------------------
    def figure(self, image_path: Path, caption: str = "", title: str = "",
               width_in: float = CONTENT_WIDTH_IN) -> int:
        path = Path(image_path)
        if not path.exists():
            LOGGER.warning("figure missing: %s", path)
            return self._figure_no

        self._figure_no += 1
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.keep_with_next = True
        paragraph.add_run().add_picture(str(path), width=Inches(width_in))

        label = f"Figure {self._figure_no}"
        if title:
            label += f". {title}"
        caption_par = self.doc.add_paragraph()
        caption_par.alignment = WD_ALIGN_PARAGRAPH.LEFT
        caption_par.paragraph_format.space_after = Pt(10)
        run = caption_par.add_run(label + ("  " if caption else ""))
        run.font.size = Pt(8.5)
        run.font.bold = True
        run.font.color.rgb = ACCENT_DARK
        if caption:
            run = caption_par.add_run(str(caption))
            run.font.size = Pt(8.5)
            run.font.color.rgb = INK_SECONDARY
        return self._figure_no

    # -- tables ---------------------------------------------------------
    def table(self, frame: pd.DataFrame, title: str = "", include_index: bool = True,
              index_label: str = "", max_rows: int = 30, note: str = "",
              highlight_sign_cols: Sequence[str] = ()) -> int:
        if frame is None or len(frame) == 0:
            return self._table_no

        data = frame.head(max_rows).copy()
        truncated = len(frame) > max_rows

        if include_index:
            data = data.reset_index()
            first = data.columns[0]
            data = data.rename(columns={first: index_label or _pretty(str(first))})

        self._table_no += 1
        if title:
            caption = self.doc.add_paragraph()
            caption.paragraph_format.space_before = Pt(8)
            caption.paragraph_format.space_after = Pt(3)
            caption.paragraph_format.keep_with_next = True
            run = caption.add_run(f"Table {self._table_no}. ")
            run.font.size = Pt(8.5)
            run.font.bold = True
            run.font.color.rgb = ACCENT_DARK
            run = caption.add_run(str(title))
            run.font.size = Pt(8.5)
            run.font.bold = True
            run.font.color.rgb = INK_SECONDARY

        columns = [_pretty(str(c)) for c in data.columns]
        table = self.doc.add_table(rows=1, cols=len(columns))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True

        sign_indices = {
            i for i, c in enumerate(data.columns)
            if str(c) in set(highlight_sign_cols) or _looks_like_change(str(c))
        }

        header = table.rows[0]
        for i, name in enumerate(columns):
            cell = header.cells[i]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if i else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(name)
            run.font.size = Pt(8.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            _shade(cell, HEADER_FILL)
            _cell_margins(cell)

        for row_no, (_, row) in enumerate(data.iterrows()):
            cells = table.add_row().cells
            for i, value in enumerate(row):
                cell = cells[i]
                cell.text = ""
                paragraph = cell.paragraphs[0]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT if i and _is_number(value) else (
                    WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
                )
                paragraph.paragraph_format.space_after = Pt(0)
                run = paragraph.add_run(_cell_text(value))
                run.font.size = Pt(8.5)
                run.font.color.rgb = INK_PRIMARY
                if i == 0:
                    run.font.bold = True
                if i in sign_indices and _is_number(value):
                    number = float(value)
                    run.font.color.rgb = POSITIVE if number >= 0 else NEGATIVE
                if row_no % 2 == 1:
                    _shade(cell, BAND_FILL)
                _cell_margins(cell)

        _table_borders(table)

        if truncated or note:
            parts = []
            if truncated:
                parts.append(f"Showing {max_rows} of {len(frame)} rows.")
            if note:
                parts.append(note)
            self.para(" ".join(parts), size=8, color=INK_MUTED, italic=True, align="left", space_after=10)
        else:
            self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return self._table_no

    # -- document furniture ---------------------------------------------
    def cover(self, meta: dict[str, str]) -> None:
        for _ in range(3):
            self.doc.add_paragraph()

        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = paragraph.add_run("HOUSING MARKET RESEARCH")
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = ACCENT

        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(4)
        run = paragraph.add_run(self.title)
        run.font.size = Pt(28)
        run.font.bold = True
        run.font.color.rgb = INK_PRIMARY

        if self.subtitle:
            paragraph = self.doc.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(18)
            run = paragraph.add_run(self.subtitle)
            run.font.size = Pt(13)
            run.font.color.rgb = INK_SECONDARY

        self._accent_rule()

        table = self.doc.add_table(rows=0, cols=2)
        for key, value in meta.items():
            if not value:
                continue
            cells = table.add_row().cells
            cells[0].text = ""
            run = cells[0].paragraphs[0].add_run(str(key))
            run.font.size = Pt(9)
            run.font.bold = True
            run.font.color.rgb = INK_MUTED
            cells[1].text = ""
            run = cells[1].paragraphs[0].add_run(str(value))
            run.font.size = Pt(9)
            run.font.color.rgb = INK_SECONDARY
            _cell_margins(cells[0], top=30, bottom=30)
            _cell_margins(cells[1], top=30, bottom=30)
        try:
            table.columns[0].width = Inches(1.7)
            table.columns[1].width = Inches(4.7)
        except Exception:  # pragma: no cover
            pass

        self.page_break()

    def toc(self, heading: str = "Contents") -> None:
        self.doc.add_heading(heading, level=1)
        self._accent_rule()
        paragraph = self.doc.add_paragraph()
        run = paragraph.add_run()
        _field(run, r'TOC \o "1-2" \h \z \u')
        self.para(
            "Right-click the table of contents in Word and choose “Update Field” to populate page numbers.",
            size=8, color=INK_MUTED, italic=True, align="left",
        )
        self.page_break()

    def page_break(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def footer(self, text: str) -> None:
        for section in self.doc.sections:
            paragraph = section.footer.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(text + "   |   ")
            run.font.size = Pt(8)
            run.font.color.rgb = INK_MUTED
            run = paragraph.add_run()
            _field(run, "PAGE")
            run.font.size = Pt(8)
            run.font.color.rgb = INK_MUTED
            run = paragraph.add_run(" / ")
            run.font.size = Pt(8)
            run.font.color.rgb = INK_MUTED
            run = paragraph.add_run()
            _field(run, "NUMPAGES")
            run.font.size = Pt(8)
            run.font.color.rgb = INK_MUTED

    def hyperlink(self, text: str, url: str, prefix: str = "") -> None:
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if prefix:
            run = paragraph.add_run(prefix)
            run.font.size = Pt(9)
            run.font.color.rgb = INK_SECONDARY
        _add_hyperlink(paragraph, text or url, url)

    def save(self, path: Path) -> Path:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(path))
        return path


# ---------------------------------------------------------------------------
# low-level docx helpers
# ---------------------------------------------------------------------------
def _field(run, instruction: str) -> None:
    """Insert a Word field (TOC, PAGE, NUMPAGES) into a run."""
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = " "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, placeholder, end):
        run._r.append(element)


def _add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "184f95")
    rPr.append(color)

    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rPr.append(underline)

    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")  # half-points -> 9pt
    rPr.append(size)

    run.append(rPr)
    text_element = OxmlElement("w:t")
    text_element.set(qn("xml:space"), "preserve")
    text_element.text = text
    run.append(text_element)
    link.append(run)
    paragraph._p.append(link)


def _shade(cell, hex_fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), hex_fill)
    cell._tc.get_or_add_tcPr().append(shading)


def _cell_margins(cell, top: int = 60, bottom: int = 60, left: int = 90, right: int = 90) -> None:
    """Cell padding in twentieths of a point."""
    tcPr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for name, value in (("top", top), ("start", left), ("bottom", bottom), ("end", right)):
        element = OxmlElement(f"w:{name}")
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    tcPr.append(margins)


def _table_borders(table, color: str = RULE_FILL, size: int = 4,
                   edges: Sequence[str] = ("top", "bottom", "insideH")) -> None:
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        if edge in edges:
            element.set(qn("w:val"), "single")
            element.set(qn("w:sz"), str(size))
            element.set(qn("w:space"), "0")
            element.set(qn("w:color"), color)
        else:
            element.set(qn("w:val"), "none")
            element.set(qn("w:sz"), "0")
        borders.append(element)
    tblPr.append(borders)


def _pretty(name: str) -> str:
    text = str(name).replace("_", " ").strip()
    replacements = {
        "yoy pct": "YoY %", "pct": "%", "cagr pct": "CAGR %", "change pct": "Change %",
        "share pct": "Share %", "r squared": "R²", "p value": "p-value",
    }
    lowered = text.lower()
    for key, value in replacements.items():
        if lowered == key:
            return value
    if lowered.endswith(" pct"):
        text = text[:-4] + " %"
    return text[:1].upper() + text[1:]


def _is_number(value: Any) -> bool:
    if isinstance(value, bool):
        return False
    try:
        number = float(value)
    except (TypeError, ValueError):
        return False
    return number == number


def _looks_like_change(name: str) -> bool:
    lowered = name.lower()
    return any(token in lowered for token in ("yoy", "change", "growth", "contribution", "deviation"))


def _cell_text(value: Any) -> str:
    if value is None or (isinstance(value, float) and value != value):
        return "—"
    if isinstance(value, (pd.Timestamp, date)):
        return pd.Timestamp(value).strftime("%Y-%m-%d")
    if isinstance(value, bool):
        return "yes" if value else "no"
    if _is_number(value):
        number = float(value)
        if number == int(number) and abs(number) < 1e15:
            return f"{int(number):,}"
        magnitude = abs(number)
        if magnitude >= 1e9:
            return f"{number/1e9:,.2f}B"
        if magnitude >= 1e6:
            return f"{number/1e6:,.2f}M"
        if magnitude >= 1000:
            return f"{number:,.1f}"
        return f"{number:,.2f}"
    text = str(value)
    return text if len(text) <= 220 else text[:217] + "…"
